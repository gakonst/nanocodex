#!/usr/bin/env python3
"""Offline Hand workflows. Pass an output directory to retain the artifacts."""
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile


def run(*args):
    subprocess.run(args, check=True, timeout=180)


def check(output):
    output.mkdir(parents=True, exist_ok=True)
    os.chdir(output)
    # All caches stay writable with the Docker Hand's read-only root filesystem.
    os.environ.update(PWD=str(output), MPLCONFIGDIR=str(output / "matplotlib-cache"),
                      GOCACHE=str(output / "go-cache"),
                      XDG_CACHE_HOME=str(output / "cache"))
    for command in ("git", "gh", "rg", "fd", "jq", "cc", "cmake", "node", "pnpm",
                    "uv", "go", "rustc", "cargo", "blender", "ffmpeg", "inkscape",
                    "libreoffice", "pandoc", "pdftoppm", "qpdf", "tesseract", "dot"):
        assert shutil.which(command), f"Missing executable: {command}"
    assert shutil.which("chromium") or shutil.which("google-chrome"), "Missing browser"
    import numpy as np
    import pandas as pd
    from scipy.linalg import solve
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from docx import Document
    from pptx import Presentation
    from reportlab.pdfgen import canvas
    from pypdf import PdfReader
    from PIL import Image
    from openpyxl import load_workbook
    import sympy

    assert np.allclose(solve([[2, 0], [0, 4]], [4, 8]), [2, 2])
    assert sympy.integrate(sympy.Symbol("x"), (sympy.Symbol("x"), 0, 2)) == 2
    frame = pd.DataFrame({"x": [1, 2, 3], "y": [1, 4, 9]})
    frame.to_excel("data.xlsx", index=False)
    assert load_workbook("data.xlsx").active["B4"].value == 9
    plt.plot(frame.x, frame.y)
    plt.savefig("plot.png")
    doc = Document()
    doc.add_heading("Hand toolkit", 0)
    doc.add_picture("plot.png")
    doc.save("report.docx")
    assert Document("report.docx").paragraphs[0].text == "Hand toolkit"
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[5]).shapes.title.text = "Hand toolkit"
    deck.save("slides.pptx")
    assert len(Presentation("slides.pptx").slides) == 1
    pdf = canvas.Canvas("report.pdf")
    pdf.drawString(72, 720, "Hand toolkit offline check")
    pdf.save()
    assert "Hand toolkit" in PdfReader("report.pdf").pages[0].extract_text()
    run("pdftoppm", "-singlefile", "-scale-to", "256", "-png", "report.pdf", "pdf-page")
    Image.open("pdf-page.png").verify()
    run("libreoffice", "-env:UserInstallation=file://" + str(output / "office-profile"),
        "--headless", "--convert-to", "pdf", "--outdir", "office", "report.docx")
    assert len(PdfReader("office/report.pdf").pages) >= 1
    Path("graph.dot").write_text("digraph { code -> artifacts }")
    run("dot", "-Tsvg", "graph.dot", "-o", "graph.svg")
    run("inkscape", "graph.svg", "--export-type=png", "--export-filename=graph.png")
    Image.open("graph.png").verify()
    Path("notes.md").write_text("# Toolkit\nOffline document conversion.\n")
    run("pandoc", "notes.md", "-o", "notes.docx")
    assert Document("notes.docx").paragraphs[0].text == "Toolkit"
    Path("page.html").write_text("<html><body><h1>Hand toolkit</h1></body></html>")
    # The surrounding Hand supplies the process isolation; the browser cannot
    # create its own user namespace under Docker's no-new-privileges policy.
    browser = shutil.which("chromium") or shutil.which("google-chrome")
    run(browser, "--headless", "--no-sandbox", "--disable-dev-shm-usage",
        "--no-first-run", "--no-default-browser-check",
        "--screenshot=" + str(output / "browser.png"),
        (output / "page.html").as_uri())
    Image.open("browser.png").verify()
    run("ffmpeg", "-v", "error", "-f", "lavfi", "-i", "color=c=blue:s=64x64:d=0.2",
        "-y", "clip.mp4")
    Path("render.py").write_text("import bpy\n"
        "bpy.context.scene.render.engine='CYCLES'\n"
        "bpy.context.scene.cycles.device='CPU'\n"
        "bpy.context.scene.cycles.samples=1\n"
        # Debian's Blender omits OpenImageDenoise; CPU rendering still works.
        "bpy.context.scene.cycles.use_denoising=False\n"
        "bpy.context.scene.render.resolution_x=32\n"
        "bpy.context.scene.render.resolution_y=32\n"
        "bpy.context.scene.render.resolution_percentage=100\n"
        f"bpy.context.scene.render.filepath={str(output / 'blender.png')!r}\n"
        "bpy.ops.render.render(write_still=True)\n")
    run("blender", "--background", "--factory-startup", "--threads", "1", "--python-exit-code", "1", "--python", str(output / "render.py"))
    Image.open("blender.png").verify()
    Path("hello.c").write_text('int main(void) { return 0; }\n')
    run("cc", "hello.c", "-o", "hello-c")
    run("./hello-c")
    Path("hello.rs").write_text('fn main() { println!("rust ok"); }\n')
    run("rustc", "hello.rs", "-o", "hello-rust")
    run("./hello-rust")
    Path("hello.go").write_text('package main\nfunc main() {}\n')
    run("go", "build", "-o", "hello-go", "hello.go")
    run("./hello-go")
    run("node", "-e", "if (2 + 2 !== 4) process.exit(1)")
    print(f"Hand toolkit passed: documents, PDF rendering, science, video, Blender CPU, C, Rust, Go, Node. Artifacts: {output}")


if __name__ == "__main__":
    if len(sys.argv) > 1:
        check(Path(sys.argv[1]).resolve())
    else:
        with tempfile.TemporaryDirectory(prefix="hand-toolkit-", dir=Path.cwd()) as temporary:
            check(Path(temporary))
